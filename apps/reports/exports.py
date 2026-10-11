import csv
import io

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _filename(report, ext):
    return f"{report['key']}_{timezone.localdate():%Y%m%d}.{ext}"


def export_csv(report):
    buf = io.StringIO()
    buf.write("﻿")  # BOM ليفتح Excel العربية بشكل صحيح
    writer = csv.writer(buf)
    writer.writerow([report["title"]])
    writer.writerow(report["columns"])
    writer.writerows(report["rows"])
    resp = HttpResponse(buf.getvalue(), content_type="text/csv; charset=utf-8")
    resp["Content-Disposition"] = f'attachment; filename="{_filename(report, "csv")}"'
    return resp


def export_excel(report):
    wb = Workbook()
    ws = wb.active
    ws.title = "التقرير"
    ws.sheet_view.rightToLeft = True
    ws.append([report["title"]])
    ws["A1"].font = Font(bold=True, size=14, color="CF6028")
    ws.append([f"تاريخ الإصدار: {timezone.localtime():%Y-%m-%d %H:%M}"])
    ws.append([])
    ws.append(report["columns"])
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="CF6028")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in report["rows"]:
        ws.append(row)
    for i, col in enumerate(report["columns"], 1):
        width = max([len(str(col))] + [len(str(r[i - 1])) for r in report["rows"][:200] if len(r) >= i])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{_filename(report, "xlsx")}"'
    return resp
