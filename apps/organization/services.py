"""خدمات الموظفين: الحفظ مع حساب الدخول، والاستيراد الجماعي من Excel."""
import io
import secrets
import string
from datetime import date, datetime

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from apps.accounts.models import Role, User, normalize_phone
from apps.audit import services as audit

from .models import Employee, EmployeeUnit, Unit


def temp_password():
    alphabet = string.ascii_letters + string.digits
    core = "".join(secrets.choice(alphabet) for _ in range(8))
    return f"Eh{core}#{secrets.randbelow(90) + 10}"


@transaction.atomic
def save_employee(form, actor):
    """يحفظ الموظف ووحداته ورئاسة الوحدة وحساب الدخول من EmployeeForm."""
    d = form.cleaned_data
    is_new = not form.instance.pk
    emp = form.save()
    units = list(d["units"])
    primary = d["primary_unit"]
    EmployeeUnit.objects.filter(employee=emp).exclude(unit__in=units).delete()
    for u in units:
        EmployeeUnit.objects.update_or_create(employee=emp, unit=u, defaults={"is_primary": u == primary})
    if d.get("is_unit_head"):
        Unit.objects.filter(pk=primary.pk).update(head=emp)
    else:
        Unit.objects.filter(head=emp, pk=primary.pk).update(head=None)

    user = emp.user
    if d.get("has_account"):
        if user is None:
            user = User.objects.create_user(d["email"], d["password"], full_name=emp.full_name, phone=d.get("phone"),
                                            role=d["role"], is_active=d.get("account_active", True),
                                            must_change_password=True)
            emp.user = user
            emp.save(update_fields=["user"])
            audit.log("user.created", user=actor, obj=user, description=f"حساب دخول للموظف {emp.full_name}")
        else:
            user.email, user.phone, user.role = d["email"], d.get("phone"), d["role"]
            user.full_name = emp.full_name
            user.is_active = d.get("account_active", True)
            if d.get("password"):
                user.set_password(d["password"])
                user.must_change_password = True
            user.save()
            audit.log("user.updated", user=actor, obj=user, description=f"تحديث حساب {emp.full_name}")
    elif user is not None:
        user.is_active = False  # إلغاء الحساب = تعطيله (لا يُحذف حفاظًا على سجل النشاط)
        user.save(update_fields=["is_active"])
    audit.log("employee.created" if is_new else "employee.updated", user=actor, obj=emp, description=emp.full_name)
    return emp


# ===== الاستيراد من Excel =====

COLUMNS = [
    ("full_name", "الاسم الكامل *"),
    ("employee_no", "رقم الموظف *"),
    ("job_title", "المسمى الوظيفي"),
    ("unit", "الوحدة الأساسية * (الاسم أو الرمز)"),
    ("other_units", "وحدات إضافية (مفصولة بفاصلة)"),
    ("is_head", "رئيس الوحدة؟ (نعم/لا)"),
    ("manager_no", "رقم الموظف للمدير المباشر"),
    ("email", "البريد الإلكتروني (لإنشاء حساب)"),
    ("phone", "رقم الجوال"),
    ("role", "الدور"),
    ("join_date", "تاريخ الانضمام (YYYY-MM-DD)"),
    ("password", "كلمة مرور مؤقتة (اختياري)"),
]
ROLE_ALIASES = {
    "مسؤول النظام": "admin", "المدير التنفيذي": "executive_director", "مساعد المدير التنفيذي": "ed_assistant",
    "رئيس الوحدة": "unit_head", "رئيس وحدة": "unit_head", "موظف": "employee", "مستخدم للعرض فقط": "viewer",
    "عرض فقط": "viewer",
}
YES = {"نعم", "yes", "y", "1", "true", "x", "✓"}


def template_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "الموظفون"
    ws.sheet_view.rightToLeft = True
    ws.append([label for _, label in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="CF6028")
    ws.append(["أحمد بن سالم الجهني", "EH-1001", "مطور أنظمة", "وحدة التقنية والتحول الرقمي", "", "لا", "EH-1000",
               "ahmad@example.org", "0500000001", "موظف", "2024-01-15", ""])
    for i in range(1, len(COLUMNS) + 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 26
    roles = DataValidation(type="list", formula1='"موظف,رئيس الوحدة,المدير التنفيذي,مساعد المدير التنفيذي,مسؤول النظام,مستخدم للعرض فقط"')
    yesno = DataValidation(type="list", formula1='"نعم,لا"')
    ws.add_data_validation(roles)
    ws.add_data_validation(yesno)
    roles.add("J2:J1000")
    yesno.add("F2:F1000")
    units = wb.create_sheet("الوحدات المتاحة")
    units.sheet_view.rightToLeft = True
    units.append(["اسم الوحدة", "الرمز"])
    for u in Unit.objects.filter(is_active=True):
        units.append([u.name, u.code or ""])
    units.column_dimensions["A"].width = 40
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _find_unit(value, units_by_key):
    return units_by_key.get(value.strip().lower()) if value else None


def _parse_date(v):
    if not v:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return date.fromisoformat(str(v)[:10])


@transaction.atomic
def import_employees(file, actor, create_missing_units=False):
    """
    يتحقق من كل الصفوف أولًا؛ إن وُجد أي خطأ لا يُستورد شيء.
    يعيد (المنشأون, الأخطاء) حيث المنشأون قائمة تتضمن كلمات المرور المؤقتة لعرضها للمسؤول مرة واحدة.
    """
    try:
        wb = load_workbook(file, read_only=True, data_only=True)
    except Exception:
        raise ValidationError("تعذر قراءة الملف. استخدم القالب بصيغة .xlsx")
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    keys = [k for k, _ in COLUMNS]
    roles = {r.code: r for r in Role.objects.all()}
    units_by_key = {}
    for u in Unit.objects.all():
        units_by_key[u.name.lower()] = u
        if u.code:
            units_by_key[u.code.lower()] = u
    existing_nos = set(Employee.objects.values_list("employee_no", flat=True))
    existing_emails = {e.lower() for e in User.objects.values_list("email", flat=True)}
    existing_phones = set(User.objects.exclude(phone=None).values_list("phone", flat=True))

    parsed, errors = [], []
    seen_nos, seen_emails = set(), set()
    for idx, raw in enumerate(rows, start=2):
        values = dict(zip(keys, [_cell(v) for v in (list(raw) + [None] * len(keys))[:len(keys)]]))
        if not any(values.values()):
            continue
        err = []
        if not values["full_name"]:
            err.append("الاسم مطلوب")
        no = str(values["employee_no"])
        if not no:
            err.append("رقم الموظف مطلوب")
        elif no in existing_nos or no in seen_nos:
            err.append(f"رقم الموظف {no} مكرر")
        seen_nos.add(no)
        unit = _find_unit(str(values["unit"]), units_by_key)
        if not unit:
            if values["unit"] and create_missing_units:
                unit = Unit.objects.create(name=str(values["unit"]))
                units_by_key[unit.name.lower()] = unit
            else:
                err.append(f"الوحدة «{values['unit']}» غير موجودة")
        extra_units = []
        for name in [x for x in str(values["other_units"]).replace("،", ",").split(",") if x.strip()]:
            u = _find_unit(name, units_by_key)
            if u:
                extra_units.append(u)
            else:
                err.append(f"الوحدة الإضافية «{name.strip()}» غير موجودة")
        email = str(values["email"]).lower()
        phone = normalize_phone(str(values["phone"])) if values["phone"] else None
        if email:
            if "@" not in email:
                err.append("بريد غير صالح")
            elif email in existing_emails or email in seen_emails:
                err.append(f"البريد {email} مستخدم")
            seen_emails.add(email)
            if phone and phone in existing_phones:
                err.append(f"الجوال {phone} مستخدم")
        role_value = str(values["role"]) or "موظف"
        role_code = ROLE_ALIASES.get(role_value, role_value)
        if email and role_code not in roles:
            err.append(f"الدور «{role_value}» غير معروف")
        try:
            join = _parse_date(values["join_date"])
        except ValueError:
            err.append("تاريخ الانضمام غير صالح")
            join = None
        if err:
            errors.append((idx, values.get("full_name") or "—", err))
            continue
        parsed.append({**values, "employee_no": no, "unit_obj": unit, "extra_units": extra_units, "email": email,
                       "phone": phone, "role_code": role_code, "join": join})
    if errors:
        transaction.set_rollback(True)
        return [], errors

    created, by_no = [], {}
    for p in parsed:
        emp = Employee.objects.create(employee_no=p["employee_no"], full_name=p["full_name"], job_title=p["job_title"],
                                      join_date=p["join"])
        EmployeeUnit.objects.create(employee=emp, unit=p["unit_obj"], is_primary=True)
        for u in p["extra_units"]:
            if u != p["unit_obj"]:
                EmployeeUnit.objects.get_or_create(employee=emp, unit=u, defaults={"is_primary": False})
        if str(p["is_head"]).lower() in YES:
            Unit.objects.filter(pk=p["unit_obj"].pk).update(head=emp)
        password = ""
        if p["email"]:
            password = p["password"] or temp_password()
            user = User.objects.create_user(p["email"], password, full_name=p["full_name"], phone=p["phone"],
                                            role=roles[p["role_code"]], must_change_password=True)
            emp.user = user
            emp.save(update_fields=["user"])
        by_no[p["employee_no"]] = emp
        created.append({"name": emp.full_name, "no": emp.employee_no, "unit": p["unit_obj"].name, "email": p["email"],
                        "role": roles[p["role_code"]].name if p["email"] else "—", "password": password,
                        "manager_no": p["manager_no"]})
    for c in created:
        if c["manager_no"]:
            mgr = by_no.get(str(c["manager_no"])) or Employee.objects.filter(employee_no=str(c["manager_no"])).first()
            if mgr:
                Employee.objects.filter(employee_no=c["no"]).update(manager=mgr)
    audit.log("employee.created", user=actor, description=f"استيراد {len(created)} موظف من Excel")
    return created, []
